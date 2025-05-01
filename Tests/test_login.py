import time
import pytest

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support import expected_conditions as EC
import openpyxl

from openpyxl import load_workbook

def read_data_from_Excel(file_path,sheet_name):
    workbook = load_workbook(file_path)
    sheet = workbook[sheet_name]
    data = []
    for row in sheet.iter_rows(values_only=True):
        data.append(row)


    workbook.close()
    return data
@pytest.mark.parametrize("username,password,result",read_data_from_Excel("projectexcel.xlsx","Sheet1")[1:])
def test_driver(username,password,result):
    print(username,password,result)
    driver = webdriver.Chrome()
    driver.get("https://www.saucedemo.com")
    driver.find_element(By.ID,"user-name").send_keys(username)
    driver.find_element(By.ID,"password").send_keys(password)
    driver.find_element(By.ID,"login-button").click()
    time.sleep(2)
    print(driver.get_cookies())
    driver.maximize_window()
    driver.quit()



